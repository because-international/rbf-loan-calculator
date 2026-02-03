#!/usr/bin/env python3
"""Create XLSX RBF Calculator for Google Sheets Import - Improved Version"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

wb = Workbook()
ws = wb.active
ws.title = "RBF Calculator"

# Define styles
header_font = Font(bold=True, size=16, color="FFFFFF")
section_font = Font(bold=True, size=12)
label_font = Font(size=11)
input_font = Font(size=11, color="0000FF")  # Blue for inputs
calculated_font = Font(size=11, bold=True, color="006400")  # Green for calculated
header_fill = PatternFill("solid", fgColor="2563EB")
section_fill = PatternFill("solid", fgColor="E5E7EB")
calculated_fill = PatternFill("solid", fgColor="DCFCE7")
derived_fill = PatternFill("solid", fgColor="F3F4F6")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

# Column widths
ws.column_dimensions['A'].width = 25
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 14
# Hidden columns
ws.column_dimensions['E'].width = 3
ws.column_dimensions['F'].width = 15  # User inputs
ws.column_dimensions['G'].width = 15  # Calculated values
ws.column_dimensions['H'].width = 15  # Alternative Repay Obl

# === HEADER ===
ws.merge_cells('A1:D1')
ws['A1'] = "REVENUE BASED FINANCE CALCULATOR"
ws['A1'].font = header_font
ws['A1'].fill = header_fill
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 35

# === SOLVE FOR DROPDOWN ===
ws['A3'] = "SOLVE FOR:"
ws['A3'].font = section_font
ws['B3'] = "Repayment Period"
ws['B3'].font = Font(size=12, bold=True)
ws['B3'].fill = PatternFill("solid", fgColor="DBEAFE")
ws['B3'].border = thin_border

# Create dropdown validation
dv = DataValidation(
    type="list",
    formula1='"Factor Rate,Amount Received,Revenue Share Rate,Repayment Period,Annual Revenue,Profit Margin"',
    allow_blank=False
)
dv.error = "Please select from the list"
dv.errorTitle = "Invalid Selection"
ws.add_data_validation(dv)
dv.add('B3')

# === VARIABLE TABLE HEADERS ===
headers = [("Variable", "A6"), ("Value", "B6"), ("Unit", "C6"), ("Status", "D6")]
for text, cell in headers:
    ws[cell] = text
    ws[cell].font = section_font
    ws[cell].fill = section_fill
    ws[cell].border = thin_border

# === HIDDEN AREA HEADERS ===
ws['E6'] = ""
ws['F6'] = "Input"
ws['G6'] = "Calc"
ws['H6'] = "AltRepObl"
for cell in ['F6', 'G6', 'H6']:
    ws[cell].font = Font(size=9, bold=True)
    ws[cell].fill = section_fill

# === VARIABLE DATA ===
# Row assignments
rows = {
    'factorRate': 7,
    'amountReceived': 8,
    'revenueShareRate': 9,
    'repaymentPeriod': 10,
    'annualRevenue': 11,
    'profitMargin': 12,
}

# Labels and units
labels = {
    'factorRate': ('Factor Rate', 'x'),
    'amountReceived': ('Amount Received', '$'),
    'revenueShareRate': ('Revenue Share Rate', '%'),
    'repaymentPeriod': ('Repayment Period', 'mo'),
    'annualRevenue': ('Annual Revenue', '$'),
    'profitMargin': ('Profit Margin', '%'),
}

# Default input values (Column F - user editable)
defaults = {
    'factorRate': 1.5,
    'amountReceived': 5000,
    'revenueShareRate': 5,
    'repaymentPeriod': 24,
    'annualRevenue': 22000,
    'profitMargin': 16,
}

# Set up labels, units, and defaults
for var, row in rows.items():
    label, unit = labels[var]
    ws[f'A{row}'] = label
    ws[f'A{row}'].font = label_font
    ws[f'A{row}'].border = thin_border
    ws[f'C{row}'] = unit
    ws[f'C{row}'].font = label_font
    ws[f'C{row}'].alignment = Alignment(horizontal='center')
    ws[f'C{row}'].border = thin_border
    # Default input value in column F
    ws[f'F{row}'] = defaults[var]
    ws[f'F{row}'].font = input_font

# === SOLVE-FOR FORMULAS (Column G) ===
# These formulas calculate each variable when it's the solve-for target

# Factor Rate: (monthlyPayment * repaymentPeriod) / amountReceived
# monthlyPayment = (annualRevenue/12) * (revenueShareRate/100)
# When solving for factorRate, use inputs for other variables directly
ws['G7'] = '=IF(OR(F10=0,F8=0),0,((F11/12)*(F9/100)*F10)/F8)'

# Amount Received: repaymentObligation / factorRate
# Uses H8 (alternative RepayObl) to avoid circular ref
ws['G8'] = '=IF(OR(F7=0,H8=0),0,H8/F7)'

# Revenue Share Rate: (requiredMonthlyPayment / monthlyRevenue) * 100
# repaymentObligation = amountReceived * factorRate
# requiredMonthlyPayment = repaymentObligation / repaymentPeriod
ws['G9'] = '=IF(OR(F10=0,F11=0),0,(((F8*F7)/F10)/(F11/12))*100)'

# Repayment Period: repaymentObligation / monthlyPayment
ws['G10'] = '=IF(((F11/12)*(F9/100))=0,0,(F8*F7)/((F11/12)*(F9/100)))'

# Annual Revenue: solve from repaymentObligation, period, and share rate
ws['G11'] = '=IF(OR(F9=0,F10=0),0,(((F8*F7)/F10)/(F9/100))*12)'

# Profit Margin: passthrough (no calculation possible)
ws['G12'] = '=F12'

# === ALTERNATIVE REPAYMENT OBLIGATION (Column H) ===
# Used when solving for Amount Received to avoid circular dependency
# H8 = monthlyPayment * repaymentPeriod (doesn't depend on amountReceived)
ws['H8'] = '=(F11/12)*(F9/100)*F10'
ws['H8'].number_format = '$#,##0.00'

# === VISIBLE VALUE FORMULAS (Column B) ===
# IF solving for this variable, show calculated; else show input
for var, row in rows.items():
    dropdown_text = labels[var][0]
    ws[f'B{row}'] = f'=IF($B$3="{dropdown_text}",G{row},F{row})'
    ws[f'B{row}'].border = thin_border
    ws[f'B{row}'].alignment = Alignment(horizontal='right')

# === STATUS FORMULAS (Column D) ===
for var, row in rows.items():
    dropdown_text = labels[var][0]
    ws[f'D{row}'] = f'=IF($B$3="{dropdown_text}","CALCULATED","INPUT")'
    ws[f'D{row}'].border = thin_border
    ws[f'D{row}'].alignment = Alignment(horizontal='center')

# === DERIVED VALUES SECTION ===
ws['A14'] = "Derived Values"
ws['A14'].font = section_font

# Repayment Obligation = Amount Received * Factor Rate
# Use the active values from column B
ws['A15'] = "Repayment Obligation"
ws['B15'] = '=B8*B7'
ws['C15'] = "$"
ws['D15'] = "DERIVED"

# Cost of Capital = Repayment Obligation - Amount Received
ws['A16'] = "Cost of Capital"
ws['B16'] = '=B15-B8'
ws['C16'] = "$"
ws['D16'] = "DERIVED"

# Monthly Payment = (Annual Revenue / 12) * (Revenue Share Rate / 100)
ws['A17'] = "Monthly Payment"
ws['B17'] = '=(B11/12)*(B9/100)'
ws['C17'] = "$"
ws['D17'] = "DERIVED"

for row in [15, 16, 17]:
    for col in ['A', 'B', 'C', 'D']:
        ws[f'{col}{row}'].fill = derived_fill
        ws[f'{col}{row}'].border = thin_border
    ws[f'A{row}'].font = label_font
    ws[f'B{row}'].alignment = Alignment(horizontal='right')
    ws[f'C{row}'].alignment = Alignment(horizontal='center')
    ws[f'D{row}'].alignment = Alignment(horizontal='center')

# === LOAN SUMMARY SECTION ===
ws['A19'] = "Loan Summary"
ws['A19'].font = section_font

summary = [
    ("Monthly Revenue", '=B11/12', "$", 20),
    ("Repayment Period (Years)", '=B10/12', "years", 21),
    ("Effective Annual Rate", '=IF(OR(B10=0,B7<=0),0,(POWER(B7,12/B10)-1)*100)', "%", 22),
]

for name, formula, unit, row in summary:
    ws[f'A{row}'] = name
    ws[f'A{row}'].font = label_font
    ws[f'A{row}'].border = thin_border
    ws[f'B{row}'] = formula
    ws[f'B{row}'].border = thin_border
    ws[f'B{row}'].alignment = Alignment(horizontal='right')
    ws[f'C{row}'] = unit
    ws[f'C{row}'].border = thin_border
    ws[f'C{row}'].alignment = Alignment(horizontal='center')

# === DATA VALIDATION FOR INPUTS ===
validations = [
    ('F7', "decimal", "between", "0.1", "10", "Factor Rate must be 0.1-10"),
    ('F8', "decimal", "greaterThan", "0", None, "Amount must be > 0"),
    ('F9', "decimal", "between", "0.1", "100", "Share Rate must be 0.1-100%"),
    ('F10', "whole", "between", "1", "360", "Period must be 1-360 months"),
    ('F11', "decimal", "greaterThan", "0", None, "Revenue must be > 0"),
    ('F12', "decimal", "between", "0", "100", "Margin must be 0-100%"),
]

for cell, vtype, op, f1, f2, err in validations:
    if f2:
        dv_item = DataValidation(type=vtype, operator=op, formula1=f1, formula2=f2)
    else:
        dv_item = DataValidation(type=vtype, operator=op, formula1=f1)
    dv_item.error = err
    ws.add_data_validation(dv_item)
    dv_item.add(cell)

# === CONDITIONAL FORMATTING ===
green_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
red_fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')

# Green highlight for each variable row when it's calculated
for var, row in rows.items():
    ws.conditional_formatting.add(f'A{row}:D{row}', FormulaRule(
        formula=[f'$D{row}="CALCULATED"'], fill=green_fill
    ))

# Red warning for Factor Rate <= 1 when calculated
ws.conditional_formatting.add('B7', FormulaRule(
    formula=['AND($D7="CALCULATED",$B7<=1)'], fill=red_fill
))

# === NUMBER FORMATTING ===
# Factor Rate
ws['B7'].number_format = '0.00'
ws['F7'].number_format = '0.00'
ws['G7'].number_format = '0.00'

# Currency (Amount Received, Annual Revenue)
for row in [8, 11]:
    ws[f'B{row}'].number_format = '$#,##0.00'
    ws[f'F{row}'].number_format = '$#,##0.00'
    ws[f'G{row}'].number_format = '$#,##0.00'

# Percentage (Revenue Share Rate, Profit Margin)
for row in [9, 12]:
    ws[f'B{row}'].number_format = '0.00'
    ws[f'F{row}'].number_format = '0.00'
    ws[f'G{row}'].number_format = '0.00'

# Months (Repayment Period)
ws['B10'].number_format = '0.00'
ws['F10'].number_format = '0'
ws['G10'].number_format = '0.00'

# Derived values
ws['B15'].number_format = '$#,##0.00'
ws['B16'].number_format = '$#,##0.00'
ws['B17'].number_format = '$#,##0.00'

# Summary
ws['B20'].number_format = '$#,##0.00'
ws['B21'].number_format = '0.00'
ws['B22'].number_format = '0.00'

# === HIDE CALCULATION COLUMNS ===
for col in ['E', 'F', 'G', 'H']:
    ws.column_dimensions[col].hidden = True

# === INSTRUCTIONS ===
ws['A24'] = "Instructions:"
ws['A24'].font = section_font

instructions = [
    "1. Select what to calculate from the dropdown (cell B3)",
    "2. To change input values, unhide columns E-H (select E and beyond, right-click, Unhide)",
    "3. Edit values in column F (blue text = your inputs)",
    "4. The calculator automatically computes the selected variable",
    "5. Green highlight = calculated value; other values are your inputs",
]

for i, text in enumerate(instructions, start=25):
    ws[f'A{i}'] = text
    ws[f'A{i}'].font = Font(size=10, italic=True, color="666666")

# Save
output_path = '/Users/nateritter/Sites/because-loan-calculator/rbf-calculator.xlsx'
wb.save(output_path)
print(f"Created: {output_path}")
